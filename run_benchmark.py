"""
NW Digital Operations Hub - Vishaqa Support Intelligence
Benchmark Evaluation Runner

Reads the 105-case benchmark dataset, sends each case to the LIVE running
FastAPI service (must be running on http://127.0.0.1:8001), compares actual
vs expected output, and produces a pass/fail evaluation report as an Excel
file - this is the "Evaluation and monitoring" deliverable.

USAGE:
    1. Make sure the service is running in another terminal:
       python -m uvicorn apps.vishaqa-service.main:app --port 8001 --reload
    2. Place this script in the project root (same level as apps/, packages/)
    3. Run:  python run_benchmark.py
    4. Report will be saved as: Vishaqa_Evaluation_Report.xlsx
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import requests
import sys
from datetime import datetime
from collections import defaultdict

INPUT_FILE = "NW_Vishaqa_Benchmark_Dataset.xlsx"
OUTPUT_FILE = "Vishaqa_Evaluation_Report.xlsx"
BASE_URL = "http://127.0.0.1:8001"

# Step 0: Confirm the service is reachable before running all cases
try:
    health = requests.get(f"{BASE_URL}/support/health", timeout=5)
    health.raise_for_status()
    print(f"[OK] Service is up: {health.json()}")
except Exception as e:
    print(f"[ERROR] Cannot reach the service at {BASE_URL}.")
    print("Make sure it's running: python -m uvicorn apps.vishaqa-service.main:app --port 8001 --reload")
    print(f"Details: {e}")
    sys.exit(1)

# Step 1: Load the benchmark dataset
try:
    wb_in = openpyxl.load_workbook(INPUT_FILE, data_only=True)
except FileNotFoundError:
    print(f"[ERROR] Could not find '{INPUT_FILE}' in the current folder.")
    print("Copy it into the project root first (same folder as apps/, packages/).")
    sys.exit(1)

ws_in = wb_in["Benchmark Dataset"]
rows = list(ws_in.iter_rows(min_row=2, values_only=True))
print(f"[OK] Loaded {len(rows)} test cases from '{INPUT_FILE}'")

# Step 2: Run each case against the live API
results = []
passed_count = 0

for row in rows:
    case_id, group, message, expected_product, expected_severity, expected_human_str, notes, tenant_id = row
    expected_human = str(expected_human_str).strip().upper() == "TRUE"

    actual_product = None
    actual_severity = None
    actual_human = None
    error = None

    try:
        if group == "Escalation - Direct Severity":
            payload = {
                "tenant_id": tenant_id,
                "ticket_message": "Benchmark escalation test",
                "severity": expected_severity,
            }
            resp = requests.post(f"{BASE_URL}/support/escalation", json=payload, timeout=10)
            resp.raise_for_status()
            data = resp.json()
            actual_severity = data.get("severity")
            actual_human = data.get("escalation_required")
            actual_product = "N/A"
        else:
            payload = {
                "tenant_id": tenant_id,
                "message": message if message is not None else "",
            }
            resp = requests.post(f"{BASE_URL}/support/classify", json=payload, timeout=10)
            resp.raise_for_status()
            data = resp.json()
            actual_product = data.get("product")
            actual_severity = data.get("severity")
            actual_human = data.get("human_required")

    except Exception as e:
        error = str(e)

    if error:
        status = "ERROR"
    else:
        product_ok = (group == "Escalation - Direct Severity") or (actual_product == expected_product)
        severity_ok = (actual_severity == expected_severity)
        human_ok = (actual_human == expected_human)
        status = "PASS" if (product_ok and severity_ok and human_ok) else "FAIL"

    if status == "PASS":
        passed_count += 1

    results.append({
        "case_id": case_id,
        "group": group,
        "message": message,
        "expected_product": expected_product,
        "actual_product": actual_product,
        "expected_severity": expected_severity,
        "actual_severity": actual_severity,
        "expected_human_required": expected_human,
        "actual_human_required": actual_human,
        "status": status,
        "notes": notes,
        "error": error or "",
    })

    if status == "ERROR":
        print(f"{case_id}: {status}  -->  {error}")
    else:
        print(f"{case_id}: {status}")

total = len(results)
fail_count = sum(1 for r in results if r["status"] == "FAIL")
error_count = sum(1 for r in results if r["status"] == "ERROR")
accuracy = round((passed_count / total) * 100, 1) if total else 0

print(f"\n=== SUMMARY ===")
print(f"Total: {total} | Pass: {passed_count} | Fail: {fail_count} | Error: {error_count}")
print(f"Accuracy: {accuracy}%")

# Step 3: Build the evaluation report workbook
wb_out = openpyxl.Workbook()

ws_sum = wb_out.active
ws_sum.title = "Summary"
ws_sum["A1"] = "Vishaqa AI Support Intelligence - Evaluation Report"
ws_sum["A1"].font = Font(name="Arial", size=14, bold=True)
ws_sum["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
ws_sum["A2"].font = Font(name="Arial", size=10, italic=True)

ws_sum["A4"] = "Total Test Cases"
ws_sum["B4"] = total
ws_sum["A5"] = "Passed"
ws_sum["B5"] = passed_count
ws_sum["A6"] = "Failed"
ws_sum["B6"] = fail_count
ws_sum["A7"] = "Errors (API call failed)"
ws_sum["B7"] = error_count
ws_sum["A8"] = "Overall Accuracy"
ws_sum["B8"] = f"{accuracy}%"

for r in range(4, 9):
    ws_sum[f"A{r}"].font = Font(name="Arial", size=11, bold=True)
    ws_sum[f"B{r}"].font = Font(name="Arial", size=11)

ws_sum["B8"].font = Font(name="Arial", size=13, bold=True,
                          color="0B7A34" if accuracy >= 90 else ("B26A00" if accuracy >= 75 else "B00020"))

group_stats = defaultdict(lambda: {"pass": 0, "fail": 0, "error": 0, "total": 0})
for r in results:
    g = r["group"]
    group_stats[g]["total"] += 1
    if r["status"] == "PASS":
        group_stats[g]["pass"] += 1
    elif r["status"] == "FAIL":
        group_stats[g]["fail"] += 1
    else:
        group_stats[g]["error"] += 1

ws_sum["A10"] = "Breakdown by Test Group"
ws_sum["A10"].font = Font(name="Arial", size=12, bold=True)
headers = ["Test Group", "Total", "Pass", "Fail", "Error", "Accuracy"]
header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1849D6", end_color="1849D6", fill_type="solid")
for col_idx, h in enumerate(headers, start=1):
    c = ws_sum.cell(row=11, column=col_idx, value=h)
    c.font = header_font
    c.fill = header_fill

row_num = 12
for group, stats in sorted(group_stats.items()):
    grp_accuracy = round((stats["pass"] / stats["total"]) * 100, 1) if stats["total"] else 0
    ws_sum.cell(row=row_num, column=1, value=group).font = Font(name="Arial", size=10)
    ws_sum.cell(row=row_num, column=2, value=stats["total"]).font = Font(name="Arial", size=10)
    ws_sum.cell(row=row_num, column=3, value=stats["pass"]).font = Font(name="Arial", size=10)
    ws_sum.cell(row=row_num, column=4, value=stats["fail"]).font = Font(name="Arial", size=10)
    ws_sum.cell(row=row_num, column=5, value=stats["error"]).font = Font(name="Arial", size=10)
    ws_sum.cell(row=row_num, column=6, value=f"{grp_accuracy}%").font = Font(name="Arial", size=10)
    row_num += 1

ws_sum.column_dimensions["A"].width = 32
for col in "BCDEF":
    ws_sum.column_dimensions[col].width = 12

ws_det = wb_out.create_sheet("Detailed Results")
det_headers = ["Case ID", "Test Group", "Ticket Message", "Expected Product", "Actual Product",
               "Expected Severity", "Actual Severity", "Expected human_required",
               "Actual human_required", "Status", "Notes", "Error"]

for col_idx, h in enumerate(det_headers, start=1):
    c = ws_det.cell(row=1, column=col_idx, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws_det.freeze_panes = "A2"

status_fill = {
    "PASS": PatternFill(start_color="EAF8EF", end_color="EAF8EF", fill_type="solid"),
    "FAIL": PatternFill(start_color="FFEDED", end_color="FFEDED", fill_type="solid"),
    "ERROR": PatternFill(start_color="FFF5DF", end_color="FFF5DF", fill_type="solid"),
}

for i, r in enumerate(results, start=2):
    values = [
        r["case_id"], r["group"], r["message"], r["expected_product"], r["actual_product"],
        r["expected_severity"], r["actual_severity"], str(r["expected_human_required"]),
        str(r["actual_human_required"]), r["status"], r["notes"], r["error"],
    ]
    for col_idx, val in enumerate(values, start=1):
        c = ws_det.cell(row=i, column=col_idx, value=val)
        c.font = Font(name="Arial", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if col_idx == 10:
            c.fill = status_fill.get(r["status"], PatternFill())

det_widths = [10, 24, 45, 16, 16, 14, 14, 18, 18, 10, 32, 24]
for i, w in enumerate(det_widths, start=1):
    ws_det.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

wb_out.save(OUTPUT_FILE)
print(f"\n[OK] Evaluation report saved to: {OUTPUT_FILE}")